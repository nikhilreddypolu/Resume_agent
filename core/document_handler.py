import os
from docx import Document

def extract_text_from_docx(file_path: str) -> str:
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Resume file not found at: {file_path}")
    
    try:
        doc = Document(file_path)
        full_text = [para.text for para in doc.paragraphs if para.text.strip()]
        return '\n'.join(full_text)
    except Exception as e:
        raise RuntimeError(f"Failed to parse document {file_path}. Error: {str(e)}")

def save_text_to_docx(text: str, output_path: str) -> None:
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    
    doc = Document()
    
    paragraphs = text.split('\n')
    for para in paragraphs:
        cleaned_para = para.strip()
        if not cleaned_para:
            continue
            
        if cleaned_para.isupper() and len(cleaned_para) < 30:
            doc.add_heading(cleaned_para, level=2)
        else:
            doc.add_paragraph(cleaned_para)
            
    doc.save(output_path)